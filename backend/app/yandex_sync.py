"""Best-effort sync of orders/bonuses/referrals into .xlsx files on Yandex Disk,
and upload of review photos to Yandex Disk (published as public links).

Everything here is a no-op if YANDEX_DISK_TOKEN is not configured, and every
public function swallows its own errors so a sync failure never breaks the
main request flow.
"""
import io
from datetime import datetime
from typing import Optional

import requests
from openpyxl import Workbook, load_workbook

from app.config import YANDEX_DISK_TOKEN, YANDEX_REVIEWS_DIR, YANDEX_SYNC_DIR

API_BASE = "https://cloud-api.yandex.net/v1/disk/resources"

ORDERS_PATH = f"{YANDEX_SYNC_DIR}/orders.xlsx"
BONUSES_PATH = f"{YANDEX_SYNC_DIR}/bonuses.xlsx"
REFERRALS_PATH = f"{YANDEX_SYNC_DIR}/referrals.xlsx"

HEADERS = {
    ORDERS_PATH: ["id", "client_id", "product_id", "quantity", "status", "comment", "created_at"],
    BONUSES_PATH: ["id", "client_id", "amount", "reason", "created_at"],
    REFERRALS_PATH: ["id", "referrer_client_id", "referred_client_id", "status", "created_at"],
}


def _enabled() -> bool:
    return bool(YANDEX_DISK_TOKEN)


def _auth_headers() -> dict:
    return {"Authorization": f"OAuth {YANDEX_DISK_TOKEN}"}


def _ensure_dir(path: str) -> None:
    requests.put(API_BASE, headers=_auth_headers(), params={"path": path}, timeout=15)


def _download_workbook(path: str) -> Workbook:
    resp = requests.get(
        f"{API_BASE}/download", headers=_auth_headers(), params={"path": path}, timeout=15
    )
    if resp.status_code == 200:
        href = resp.json()["href"]
        file_resp = requests.get(href, timeout=15)
        file_resp.raise_for_status()
        return load_workbook(io.BytesIO(file_resp.content))

    wb = Workbook()
    ws = wb.active
    ws.append(HEADERS[path])
    return wb


def _upload_workbook(path: str, wb: Workbook) -> None:
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    resp = requests.get(
        f"{API_BASE}/upload",
        headers=_auth_headers(),
        params={"path": path, "overwrite": "true"},
        timeout=15,
    )
    resp.raise_for_status()
    href = resp.json()["href"]
    put_resp = requests.put(href, data=buffer.read(), timeout=30)
    put_resp.raise_for_status()


def _append_row(path: str, row: list) -> None:
    if not _enabled():
        return
    try:
        _ensure_dir(YANDEX_SYNC_DIR)
        wb = _download_workbook(path)
        ws = wb.active
        ws.append(row)
        _upload_workbook(path, wb)
    except Exception as exc:  # noqa: BLE001
        print(f"[yandex_sync] append_row failed for {path}: {exc}")


def sync_order(order) -> None:
    _append_row(
        ORDERS_PATH,
        [order.id, order.client_id, order.product_id, order.quantity, order.status, order.comment, str(order.created_at)],
    )


def sync_bonus_event(event) -> None:
    _append_row(
        BONUSES_PATH,
        [event.id, event.client_id, event.amount, event.reason, str(event.created_at)],
    )


def sync_referral(referral) -> None:
    _append_row(
        REFERRALS_PATH,
        [referral.id, referral.referrer_client_id, referral.referred_client_id, referral.status, str(referral.created_at)],
    )


def resync_all(orders: list, bonus_events: list, referrals: list) -> None:
    if not _enabled():
        return
    try:
        _ensure_dir(YANDEX_SYNC_DIR)

        wb = Workbook()
        ws = wb.active
        ws.append(HEADERS[ORDERS_PATH])
        for o in orders:
            ws.append([o.id, o.client_id, o.product_id, o.quantity, o.status, o.comment, str(o.created_at)])
        _upload_workbook(ORDERS_PATH, wb)

        wb = Workbook()
        ws = wb.active
        ws.append(HEADERS[BONUSES_PATH])
        for e in bonus_events:
            ws.append([e.id, e.client_id, e.amount, e.reason, str(e.created_at)])
        _upload_workbook(BONUSES_PATH, wb)

        wb = Workbook()
        ws = wb.active
        ws.append(HEADERS[REFERRALS_PATH])
        for r in referrals:
            ws.append([r.id, r.referrer_client_id, r.referred_client_id, r.status, str(r.created_at)])
        _upload_workbook(REFERRALS_PATH, wb)
    except Exception as exc:  # noqa: BLE001
        print(f"[yandex_sync] resync_all failed: {exc}")


def upload_review_photo(file_bytes: bytes, filename: str) -> Optional[str]:
    """Uploads a review photo to Yandex Disk and returns a public URL, or None."""
    if not _enabled():
        return None
    try:
        _ensure_dir(YANDEX_SYNC_DIR)
        _ensure_dir(YANDEX_REVIEWS_DIR)

        stamp = datetime.utcnow().strftime("%Y%m%d%H%M%S%f")
        disk_path = f"{YANDEX_REVIEWS_DIR}/{stamp}_{filename}"

        resp = requests.get(
            f"{API_BASE}/upload",
            headers=_auth_headers(),
            params={"path": disk_path, "overwrite": "true"},
            timeout=15,
        )
        resp.raise_for_status()
        href = resp.json()["href"]
        put_resp = requests.put(href, data=file_bytes, timeout=30)
        put_resp.raise_for_status()

        publish_resp = requests.put(
            f"{API_BASE}/publish", headers=_auth_headers(), params={"path": disk_path}, timeout=15
        )
        publish_resp.raise_for_status()

        meta_resp = requests.get(
            API_BASE, headers=_auth_headers(), params={"path": disk_path}, timeout=15
        )
        meta_resp.raise_for_status()
        return meta_resp.json().get("public_url")
    except Exception as exc:  # noqa: BLE001
        print(f"[yandex_sync] upload_review_photo failed: {exc}")
        return None
